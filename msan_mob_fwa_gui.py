import os
import sys
import json
import queue
import threading
import math
import unittest
from datetime import datetime
import numpy as np
import pandas as pd
from scipy.spatial import cKDTree
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# Optional dependencies handling
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# ==========================================
# 1. CONFIGURATION & I18N
# ==========================================

CONFIG_FILE = ".app_config.json"
SIGNATURE = "@19walid"
LOGO_FILE = "logo.png"  # optionnel : placer à côté du script / de l'exe


def resource_path(name):
    """Chemin d'une ressource, compatible script et exe PyInstaller (_MEIPASS)."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)

I18N = {
    "FR": {
        "title": "Telecom Spatial Studio Pro - MSAN & Mobile",
        "tab_workflow": " Flux de Travail ",
        "tab_settings": " Paramètres du Calcul ",
        "tab_audit": " Rapport de Validation ",
        "sec_files": " 1. Fichiers Entrants (CSV / Excel) ",
        "sec_mapping": " 2. Correspondance des Colonnes ",
        "sec_engine_cfg": " Paramètres Géospatiaux ",
        "msan_file": "Fichier MSAN / Fixe :",
        "mob_file": "Fichier Mobile / RAN :",
        "browse": "Parcourir...",
        "btn_discovery": "Détecter les Colonnes",
        "btn_load_profile": "Charger Profil",
        "btn_save_profile": "Sauvegarder Profil",
        "btn_run": "LANCER LE CALCUL",
        "btn_cancel": "ANNULER",
        "lbl_top_n": "Nombre de voisins (Top N) :",
        "lbl_cand_pool": "Taille pool candidats (k_search) :",
        "lbl_max_dist": "Distance max tolérée (km) :",
        "lbl_locale": "Format Décimal Export :",
        "lbl_kml_export": "Générer la carte Google Earth (KML) :",
        "lbl_site_name": "Nom du Site :",
        "lbl_cell_name": "Nom de Cellule :",
        "lbl_azimut": "Azimut :",
        "lbl_lon": "Longitude :",
        "lbl_lat": "Latitude :",
        "lbl_status_init": "Prêt. Sélectionnez les fichiers puis cliquez sur 'Détecter les Colonnes'.",
        "lbl_status_discovered": "Colonnes détectées ! Vérifiez les associations et lancez le calcul.",
        "lbl_status_running": "Calcul en cours... Veuillez patienter.",
        "lbl_status_cancelled": "Calcul annulé par l'utilisateur.",
        "lbl_status_success": "Traitement terminé avec succès !",
        "lbl_status_error": "Erreur pendant le traitement.",
        "warn_no_files": "Veuillez sélectionner les deux fichiers d'entrée.",
        "err_no_valid_points": "Aucune coordonnée valide n'a été trouvée dans les fichiers.",
        "lbl_out_dir": "Dossier de sortie :",
        "err_out_dir": "Impossible d'utiliser le dossier de sortie :",
    },
    "EN": {
        "title": "Telecom Spatial Studio Pro - MSAN & Mobile",
        "tab_workflow": " Workflow ",
        "tab_settings": " Engine Settings ",
        "tab_audit": " Validation Audit ",
        "sec_files": " 1. Input Files (CSV / Excel) ",
        "sec_mapping": " 2. Column Mapping ",
        "sec_engine_cfg": " Geospatial Engine Parameters ",
        "msan_file": "MSAN / Fixed File:",
        "mob_file": "Mobile / RAN File:",
        "browse": "Browse...",
        "btn_discovery": "Detect Columns",
        "btn_load_profile": "Load Profile",
        "btn_save_profile": "Save Profile",
        "btn_run": "RUN CALCULATION",
        "btn_cancel": "CANCEL",
        "lbl_top_n": "Neighbors to return (Top N):",
        "lbl_cand_pool": "Candidate pool size (k_search):",
        "lbl_max_dist": "Max distance cap (km):",
        "lbl_locale": "Export Decimal Format:",
        "lbl_kml_export": "Generate Google Earth Map (KML):",
        "lbl_site_name": "Site Name:",
        "lbl_cell_name": "Cell Name:",
        "lbl_azimut": "Azimuth:",
        "lbl_lon": "Longitude:",
        "lbl_lat": "Latitude:",
        "lbl_status_init": "Ready. Select input files and click 'Detect Columns'.",
        "lbl_status_discovered": "Columns detected! Review mapping and run calculation.",
        "lbl_status_running": "Processing data... Please wait.",
        "lbl_status_cancelled": "Calculation cancelled by user.",
        "lbl_status_success": "Processing completed successfully!",
        "lbl_status_error": "Error encountered during execution.",
        "warn_no_files": "Please select both MSAN and Mobile input files.",
        "err_no_valid_points": "No valid coordinates found in input datasets.",
        "lbl_out_dir": "Output Folder:",
        "err_out_dir": "Cannot use the output folder:",
    }
}


# ==========================================
# 2. FILE IO & ENCODING SURVIVAL ENGINE
# ==========================================

class DataIngestionEngine:
    """Handles reading CSV/Excel safely, preserving string types and autodetecting encoding."""

    @staticmethod
    def read_file(filepath):
        ext = os.path.splitext(filepath)[1].lower()
        if ext in [".xlsx", ".xls", ".xlsb"]:
            engine = "pyxlsb" if ext == ".xlsb" else None
            df = pd.read_excel(filepath, dtype=str, engine=engine)
            return df, f"Excel ({ext})"
        
        # CSV Handling with multi-encoding fallback
        encodings = ["utf-8-sig", "utf-8", "cp1252", "latin1", "iso-8859-1", "utf-16"]
        separators = [";", ",", "\t", "|"]

        for enc in encodings:
            for sep in separators:
                try:
                    df = pd.read_csv(filepath, sep=sep, encoding=enc, dtype=str, on_bad_lines="skip")
                    if len(df.columns) > 1:
                        return df, f"CSV (enc={enc}, sep='{sep}')"
                except Exception:
                    continue

        # Last resort fallback
        df = pd.read_csv(filepath, sep=None, engine="python", dtype=str, encoding="latin1")
        return df, "CSV (Fallback Python Engine)"


# ==========================================
# 3. KML MAP EXPORTER
# ==========================================
class KMLExporter:
    """
    KML layout:

      01 - MSAN Sites
      02 - MSAN Sites Name
      03 - Mobile Sites
      04 - Mobile Sites Name
      05 - Mobile Sector
      06 - MSAN to Mobile Connections TOP1
      07 - MSAN to Mobile Connections TOP2
      08 - MSAN to Mobile Connections TOP3

    Each folder is independently show/hideable in Google Earth.
    """

    @staticmethod
    def _safe_float(value):
        try:
            if pd.isna(value):
                return None
            return float(str(value).replace(",", ".").strip())
        except (ValueError, TypeError):
            return None

    @staticmethod
    def _valid_coordinate(lat, lon):
        """Return True only for finite WGS84 latitude/longitude."""
        lat = KMLExporter._safe_float(lat)
        lon = KMLExporter._safe_float(lon)

        if lat is None or lon is None:
            return False

        if not math.isfinite(lat) or not math.isfinite(lon):
            return False

        return -90.0 <= lat <= 90.0 and -180.0 <= lon <= 180.0

    @staticmethod
    def _format_distance(value):
        """Distance lisible (ex: "1.23 km") ou "N/A" si absente / invalide."""
        dist = KMLExporter._safe_float(value)
        if dist is None or not math.isfinite(dist):
            return "N/A"
        return f"{dist:.2f} km"

    @staticmethod
    def _sector_points(lat, lon, azimuth, radius_km=0.5,
                       half_width_deg=30.0):
        if not KMLExporter._valid_coordinate(lat, lon):
            return None

        try:
            lat = float(lat)
            lon = float(lon)
            azi = float(str(azimuth).replace(",", ".").strip())
        except (ValueError, TypeError):
            return None

        if not math.isfinite(azi):
            return None

        points = [(lon, lat)]
        lat_rad = math.radians(lat)

        for angle in (
            azi - half_width_deg,
            azi - half_width_deg / 2,
            azi,
            azi + half_width_deg / 2,
            azi + half_width_deg
        ):
            angle_rad = math.radians(angle % 360)
            dlat = radius_km * math.cos(angle_rad) / 111.32
            cos_lat = max(abs(math.cos(lat_rad)), 0.01)
            dlon = radius_km * math.sin(angle_rad) / (111.32 * cos_lat)
            points.append((lon + dlon, lat + dlat))

        points.append((lon, lat))
        return points

    @staticmethod
    def export(results_df, output_kml_path, mobile_df=None, mobile_cols=None):
        from xml.sax.saxutils import escape

        rank_colors = {
            1: "ff0000ff",  # Red
            2: "ff00ff00",  # Green
            3: "ffff0000",  # Blue
        }

        kml_audit = []

        def audit(layer, status, object_type, object_name, details,
                  source_row="", reason=""):
            kml_audit.append({
                "Layer": layer,
                "Status": status,
                "Object": object_type,
                "Name": object_name,
                "Source Row": source_row,
                "Reason": reason,
                "Details": details
            })

        kml = [
            '<?xml version="1.0" encoding="UTF-8"?>',
            '<kml xmlns="http://www.opengis.net/kml/2.2">',
            '  <Document>',
            '    <name>MSAN - Mobile Top 3 Spatial Map</name>',
            '',
            '    <Style id="msanPointStyle">',
            '      <IconStyle>',
            '        <scale>1.1</scale>',
            '        <Icon>',
            '          <href>http://maps.google.com/mapfiles/kml/pushpin/ylw-pushpin.png</href>',
            '        </Icon>',
            '      </IconStyle>',
            '      <LabelStyle><scale>0.0</scale></LabelStyle>',
            '    </Style>',
            '    <Style id="mobilePointStyle">',
            '      <IconStyle>',
            '        <scale>1.0</scale>',
            '        <Icon>',
            '          <href>http://maps.google.com/mapfiles/kml/pushpin/blu-pushpin.png</href>',
            '        </Icon>',
            '      </IconStyle>',
            '      <LabelStyle><scale>0.0</scale></LabelStyle>',
            '    </Style>',
            '    <Style id="cellPointStyle">',
            '      <IconStyle><scale>0.0</scale></IconStyle>',
            '      <LabelStyle><scale>0.0</scale></LabelStyle>',
            '    </Style>',
            '    <Style id="labelStyle">',
            '      <IconStyle><scale>0.0</scale></IconStyle>',
            '      <LabelStyle><scale>1.0</scale></LabelStyle>',
            '    </Style>',
            ''
        ]

        for rank, color in rank_colors.items():
            kml += [
                f'    <Style id="lineTop{rank}">',
                '      <LineStyle>',
                f'        <color>{color}</color>',
                '        <width>4</width>',
                '      </LineStyle>',
                '    </Style>',
                f'    <Style id="sectorTop{rank}">',
                '      <LineStyle>',
                f'        <color>{color}</color>',
                '        <width>2</width>',
                '      </LineStyle>',
                '      <PolyStyle><color>33000000</color></PolyStyle>',
                '    </Style>',
                f'    <Style id="distLabelTop{rank}">',
                '      <IconStyle><scale>0.0</scale></IconStyle>',
                '      <LabelStyle>',
                f'        <color>{color}</color>',
                '        <scale>0.9</scale>',
                '      </LabelStyle>',
                '    </Style>',
                ''
            ]

        # ---------------------------------------------------------
        # 01 - MSAN Sites
        # ---------------------------------------------------------
        kml += [
            '    <Folder>',
            '      <name>01 - MSAN Sites</name>',
            '      <open>0</open>'
        ]

        msan_records = []
        msan_skipped = 0

        for _, row in results_df.iterrows():
            lat = KMLExporter._safe_float(row.get("MSAN_LAT"))
            lon = KMLExporter._safe_float(row.get("MSAN_LON"))

            if not KMLExporter._valid_coordinate(lat, lon):
                msan_skipped += 1
                source_row = int(row.name) + 2
                raw_lat = row.get("MSAN_LAT")
                raw_lon = row.get("MSAN_LON")

                if lat is None or lon is None:
                    reason = "Missing or non-numeric coordinate"
                elif not math.isfinite(lat) or not math.isfinite(lon):
                    reason = "Non-finite coordinate"
                elif not (-90.0 <= lat <= 90.0):
                    reason = "Latitude out of range [-90, 90]"
                elif not (-180.0 <= lon <= 180.0):
                    reason = "Longitude out of range [-180, 180]"
                else:
                    reason = "Invalid coordinate"

                audit(
                    "01 - MSAN Sites", "SKIPPED", "Placemark",
                    str(row.get("Site_MSAN", "MSAN")),
                    f"LAT={raw_lat}; LON={raw_lon}",
                    source_row, reason
                )
                continue

            msan_id = f"{len(msan_records) + 1:03d}"
            msan_name = str(row.get("Site_MSAN", "MSAN"))

            msan_records.append({
                "id": msan_id,
                "name": msan_name,
                "lat": lat,
                "lon": lon
            })

            kml += [
                '      <Placemark>',
                f'        <name></name>',
                '        <styleUrl>#msanPointStyle</styleUrl>',
                '        <description><![CDATA[',
                f'<b>MSAN ID:</b> {escape(msan_id)}<br/>',
                f'<b>MSAN Site Name:</b> {escape(msan_name)}<br/>',
                f'<b>Latitude:</b> {lat}<br/>',
                f'<b>Longitude:</b> {lon}',
                ']]></description>',
                '        <Point>',
                f'          <coordinates>{lon},{lat},0</coordinates>',
                '        </Point>',
                '      </Placemark>'
            ]

        audit("01 - MSAN Sites", "GENERATED", "Placemark", "MSAN Sites", f"Created: {len(msan_records)}; Skipped invalid coordinates: {msan_skipped}")
        kml += ['    </Folder>', '']

        # ---------------------------------------------------------
        # 02 - MSAN Sites Name
        # ---------------------------------------------------------
        kml += [
            '    <Folder>',
            '      <name>02 - MSAN Sites Name</name>',
            '      <open>0</open>'
        ]

        for rec in msan_records:
            kml += [
                '      <Placemark>',
                f'        <name>{escape(rec["name"])}</name>',
                '        <styleUrl>#labelStyle</styleUrl>',
                '        <Point>',
                f'          <coordinates>{rec["lon"]},{rec["lat"]},0</coordinates>',
                '        </Point>',
                '      </Placemark>'
            ]

        audit("02 - MSAN Sites Name", "GENERATED", "Label", "MSAN Site Names", f"Created: {len(msan_records)}")
        kml += ['    </Folder>', '']

        # ---------------------------------------------------------
        # Build all mobile sites and all cells from the original file.
        # ---------------------------------------------------------
        mobile_records = []
        mobile_by_site = {}
        mobile_invalid = 0
        mobile_missing_site = 0

        if mobile_df is not None and mobile_cols:
            site_col = mobile_cols["mob_site"]
            cell_col = mobile_cols["mob_cell"]
            lat_col = mobile_cols["mob_lat"]
            lon_col = mobile_cols["mob_lon"]
            azi_col = mobile_cols["mob_azimut"]

            for _, mob_row in mobile_df.iterrows():
                site = mob_row.get(site_col)

                source_row = int(mob_row.name) + 2

                if pd.isna(site) or str(site).strip() == "":
                    mobile_missing_site += 1
                    audit(
                        "03 - Mobile Sites", "SKIPPED", "Mobile Site",
                        "(missing site name)",
                        f"LAT={mob_row.get(lat_col)}; LON={mob_row.get(lon_col)}",
                        source_row,
                        "Missing Mobile Site Name"
                    )
                    continue

                site_name = str(site)
                lat_raw = mob_row.get(lat_col)
                lon_raw = mob_row.get(lon_col)
                lat = KMLExporter._safe_float(lat_raw)
                lon = KMLExporter._safe_float(lon_raw)

                if not KMLExporter._valid_coordinate(lat, lon):
                    mobile_invalid += 1

                    if lat is None or lon is None:
                        reason = "Missing or non-numeric coordinate"
                    elif not math.isfinite(lat) or not math.isfinite(lon):
                        reason = "Non-finite coordinate"
                    elif not (-90.0 <= lat <= 90.0):
                        reason = "Latitude out of range [-90, 90]"
                    elif not (-180.0 <= lon <= 180.0):
                        reason = "Longitude out of range [-180, 180]"
                    else:
                        reason = "Invalid coordinate"

                    audit(
                        "03 - Mobile Sites", "SKIPPED", "Mobile Site",
                        site_name,
                        f"LAT={lat_raw}; LON={lon_raw}",
                        source_row, reason
                    )
                    continue

                if site_name not in mobile_by_site:
                    mobile_by_site[site_name] = {
                        "id": f"{len(mobile_records) + 1:03d}",
                        "name": site_name,
                        "lat": lat,
                        "lon": lon,
                        "cells": []
                    }
                    mobile_records.append(mobile_by_site[site_name])

                mobile_by_site[site_name]["cells"].append({
                    "name": str(mob_row.get(cell_col, "")),
                    "lat": lat,
                    "lon": lon,
                    "azimuth": mob_row.get(azi_col)
                })

        # ---------------------------------------------------------
        # 03 - Mobile Sites
        # ---------------------------------------------------------
        kml += [
            '    <Folder>',
            '      <name>03 - Mobile Sites</name>',
            '      <open>0</open>'
        ]

        for site in mobile_records:
            kml += [
                '      <Placemark>',
                f'        <name></name>',
                '        <styleUrl>#mobilePointStyle</styleUrl>',
                '        <description><![CDATA[',
                f'<b>Mobile Site ID:</b> {escape(site["id"])}<br/>',
                f'<b>Mobile Site Name:</b> {escape(site["name"])}<br/>',
                f'<b>Number of Cells:</b> {len(site["cells"])}<br/>',
                f'<b>Latitude:</b> {site["lat"]}<br/>',
                f'<b>Longitude:</b> {site["lon"]}',
                ']]></description>',
                '        <Point>',
                f'          <coordinates>{site["lon"]},{site["lat"]},0</coordinates>',
                '        </Point>',
                '      </Placemark>'
            ]

        audit("03 - Mobile Sites", "GENERATED", "Placemark", "Mobile Sites", f"Created: {len(mobile_records)}; Missing site name rows: {mobile_missing_site}; Skipped invalid coordinates: {mobile_invalid}")
        kml += ['    </Folder>', '']

        # ---------------------------------------------------------
        # 04 - Mobile Sites Name
        # ---------------------------------------------------------
        kml += [
            '    <Folder>',
            '      <name>04 - Mobile Sites Name</name>',
            '      <open>0</open>'
        ]

        for site in mobile_records:
            kml += [
                '      <Placemark>',
                f'        <name>{escape(site["name"])}</name>',
                '        <styleUrl>#labelStyle</styleUrl>',
                '        <Point>',
                f'          <coordinates>{site["lon"]},{site["lat"]},0</coordinates>',
                '        </Point>',
                '      </Placemark>'
            ]

        audit("04 - Mobile Sites Name", "GENERATED", "Label", "Mobile Site Names", f"Created: {len(mobile_records)}")
        kml += ['    </Folder>', '']

        # ---------------------------------------------------------
        # 05 - Mobile Sector
        # All cells, with site name and cell name.
        # ---------------------------------------------------------
        kml += [
            '    <Folder>',
            '      <name>05 - Mobile Sector</name>',
            '      <open>0</open>'
        ]

        global_cell_id = 0
        sector_count = 0
        sector_invalid = 0

        for site in mobile_records:
            for cell in site["cells"]:
                if not KMLExporter._valid_coordinate(cell["lat"], cell["lon"]):
                    sector_invalid += 1
                    continue

                global_cell_id += 1
                cell_id = f"{global_cell_id:03d}"
                cell_name = cell["name"]
                azi = cell["azimuth"]
                azi_txt = str(azi) if pd.notna(azi) else "N/A"

                kml += [
                    '      <Placemark>',
                    f'        <name>{escape(cell_id)} - {escape(cell_name)}</name>',
                    '        <styleUrl>#cellPointStyle</styleUrl>',
                    '        <description><![CDATA[',
                    f'<b>Cell ID:</b> {escape(cell_id)}<br/>',
                    f'<b>Mobile Site:</b> {escape(site["name"])}<br/>',
                    f'<b>Cell Name:</b> {escape(cell_name)}<br/>',
                    f'<b>Azimuth:</b> {escape(azi_txt)}<br/>',
                    f'<b>Latitude:</b> {cell["lat"]}<br/>',
                    f'<b>Longitude:</b> {cell["lon"]}',
                    ']]></description>',
                    '        <Point>',
                    f'          <coordinates>{cell["lon"]},{cell["lat"]},0</coordinates>',
                    '        </Point>',
                    '      </Placemark>'
                ]

                points = KMLExporter._sector_points(
                    cell["lat"], cell["lon"], azi
                )

                if not points:
                    sector_invalid += 1
                    if pd.isna(azi) or str(azi).strip() == "":
                        reason = "Missing Azimuth - sector geometry skipped"
                    else:
                        try:
                            azi_num = float(str(azi).replace(",", ".").strip())
                            reason = (
                                "Non-finite Azimuth - sector geometry skipped"
                                if not math.isfinite(azi_num)
                                else "Invalid Azimuth - sector geometry skipped"
                            )
                        except (ValueError, TypeError):
                            reason = "Non-numeric Azimuth - sector geometry skipped"

                    audit(
                        "05 - Mobile Sector", "SKIPPED", "Sector",
                        f"{site['name']} / {cell_name}",
                        f"LAT={cell['lat']}; LON={cell['lon']}; Azimuth={azi_txt}",
                        "",
                        reason
                    )

                if points:
                    coords = " ".join(
                        f"{x:.7f},{y:.7f},0" for x, y in points
                    )

                    sector_count += 1
                    kml += [
                        '      <Placemark>',
                        f'        <name>Sector {escape(cell_id)} - '
                        f'{escape(cell_name)}</name>',
                        '        <Style>',
                        '          <LineStyle>',
                        '            <color>ff00ffff</color>',
                        '            <width>2</width>',
                        '          </LineStyle>',
                        '          <PolyStyle>',
                        '            <color>2200ffff</color>',
                        '          </PolyStyle>',
                        '        </Style>',
                        '        <description><![CDATA[',
                        f'<b>Mobile Site:</b> {escape(site["name"])}<br/>',
                        f'<b>Cell:</b> {escape(cell_name)}<br/>',
                        f'<b>Azimuth:</b> {escape(azi_txt)}',
                        ']]></description>',
                        '        <Polygon>',
                        '          <tessellate>1</tessellate>',
                        '          <outerBoundaryIs>',
                        '            <LinearRing>',
                        f'              <coordinates>{coords}</coordinates>',
                        '            </LinearRing>',
                        '          </outerBoundaryIs>',
                        '        </Polygon>',
                        '      </Placemark>'
                    ]

        audit("05 - Mobile Sector", "GENERATED", "Sector/Cell", "Mobile Sectors", f"Cells: {global_cell_id}; Sector geometries: {sector_count}; Invalid cells skipped: {sector_invalid}")
        kml += ['    </Folder>', '']

        # ---------------------------------------------------------
        # 06 / 07 / 08 - Top 1 / Top 2 / Top 3 connection folders.
        # ---------------------------------------------------------
        for rank in (1, 2, 3):
            connection_count = 0
            connection_skipped = 0
            kml += [
                '    <Folder>',
                f'      <name>0{5 + rank} - MSAN to Mobile Connections TOP{rank}</name>',
                '      <open>0</open>'
            ]

            for _, row in results_df.iterrows():
                msan_lat = KMLExporter._safe_float(row.get("MSAN_LAT"))
                msan_lon = KMLExporter._safe_float(row.get("MSAN_LON"))
                mob_lat = KMLExporter._safe_float(
                    row.get(f"Top{rank}_Mobile_LAT")
                )
                mob_lon = KMLExporter._safe_float(
                    row.get(f"Top{rank}_Mobile_LON")
                )
                mob_site = row.get(f"Top{rank}_Site_Mobile")

                if (
                    not KMLExporter._valid_coordinate(msan_lat, msan_lon)
                    or not KMLExporter._valid_coordinate(mob_lat, mob_lon)
                    or pd.isna(mob_site)
                ):
                    connection_skipped += 1
                    source_row = int(row.name) + 2

                    if pd.isna(mob_site):
                        reason = f"TOP{rank} Mobile Site is missing"
                    elif not KMLExporter._valid_coordinate(msan_lat, msan_lon):
                        reason = f"Invalid MSAN coordinate: LAT={row.get('MSAN_LAT')}; LON={row.get('MSAN_LON')}"
                    elif not KMLExporter._valid_coordinate(mob_lat, mob_lon):
                        reason = f"Invalid TOP{rank} Mobile coordinate: LAT={row.get(f'Top{rank}_Mobile_LAT')}; LON={row.get(f'Top{rank}_Mobile_LON')}"
                    else:
                        reason = "Invalid connection data"

                    audit(
                        f"0{5 + rank} - MSAN to Mobile Connections TOP{rank}",
                        "SKIPPED", "LineString",
                        f"{row.get('Site_MSAN', 'MSAN')} → {mob_site if pd.notna(mob_site) else '(missing)'}",
                        f"TOP{rank} source row: {source_row}",
                        source_row, reason
                    )
                    continue

                msan_name = str(row.get("Site_MSAN", "MSAN"))
                mob_site_name = str(mob_site)

                distance = row.get(f"Top{rank}_Distance_km")
                best_cell = row.get(f"Top{rank}_Best_Cell")
                azimuth = row.get(f"Top{rank}_Azimut")
                angle_dev = row.get(f"Top{rank}_Angle_Dev")

                best_cell_txt = str(best_cell) if pd.notna(best_cell) else "N/A"
                azimuth_txt = str(azimuth) if pd.notna(azimuth) else "N/A"
                distance_txt = KMLExporter._format_distance(distance)
                dist_suffix = f" ({distance_txt})" if distance_txt != "N/A" else ""
                angle_txt = str(angle_dev) if pd.notna(angle_dev) else "N/A"

                description = (
                    f'<b>MSAN Site:</b> {escape(msan_name)}<br/>'
                    f'<b>Mobile Site:</b> {escape(mob_site_name)}<br/>'
                    f'<b>Rank:</b> TOP {rank}<br/>'
                    f'<b>Best Cell:</b> {escape(best_cell_txt)}<br/>'
                    f'<b>Azimuth:</b> {escape(azimuth_txt)}<br/>'
                    f'<b>Distance:</b> {escape(distance_txt)}<br/>'
                    f'<b>Angle Deviation:</b> {escape(angle_txt)}°'
                )

                connection_count += 1
                kml += [
                    '      <Placemark>',
                    f'        <name>{escape(msan_name)} → '
                    f'{escape(mob_site_name)}{escape(dist_suffix)}</name>',
                    f'        <styleUrl>#lineTop{rank}</styleUrl>',
                    f'        <description><![CDATA[{description}]]></description>',
                    '        <LineString>',
                    '          <tessellate>1</tessellate>',
                    '          <coordinates>',
                    f'            {msan_lon},{msan_lat},0',
                    f'            {mob_lon},{mob_lat},0',
                    '          </coordinates>',
                    '        </LineString>',
                    '      </Placemark>'
                ]

                # Étiquette de distance visible directement sur la carte,
                # placée au milieu de la ligne (même dossier TOPn).
                if dist_suffix:
                    mid_lat = (msan_lat + mob_lat) / 2.0
                    mid_lon = (msan_lon + mob_lon) / 2.0
                    kml += [
                        '      <Placemark>',
                        f'        <name>{escape(distance_txt)}</name>',
                        f'        <styleUrl>#distLabelTop{rank}</styleUrl>',
                        '        <Point>',
                        f'          <coordinates>{mid_lon:.7f},{mid_lat:.7f},0</coordinates>',
                        '        </Point>',
                        '      </Placemark>'
                    ]

            audit(f"0{5 + rank} - MSAN to Mobile Connections TOP{rank}", "GENERATED", "LineString", f"TOP{rank} Connections", f"Created: {connection_count}; Skipped: {connection_skipped}")
            kml += ['    </Folder>', '']

        kml += [
            '  </Document>',
            '</kml>'
        ]

        with open(output_kml_path, "w", encoding="utf-8") as f:
            f.write("\n".join(kml))

        audit("KML FILE", "SUCCESS", "File", output_kml_path, f"Generated successfully with {len(kml_audit)} audit entries")
        return pd.DataFrame(kml_audit)


# ==========================================
# 4. MATH & GEOSPATIAL VECTORIZED ENGINE
# ==========================================

def haversine_np(lat1, lon1, lat2, lon2):
    R = 6371.0
    lat1, lon1, lat2, lon2 = map(np.radians, [lat1, lon1, lat2, lon2])
    dlat = lat2 - lat1
    dlon = lon2 - lon1
    a = np.sin(dlat / 2.0) ** 2 + np.cos(lat1) * np.cos(lat2) * np.sin(dlon / 2.0) ** 2
    return 2 * R * np.arcsin(np.sqrt(a))

def bearing_np(lat1, lon1, lat2, lon2):
    lat1, lon1, lat2, lon2 = map(np.radians, [lat1, lon1, lat2, lon2])
    dlon = lon2 - lon1
    y = np.sin(dlon) * np.cos(lat2)
    x = np.cos(lat1) * np.sin(lat2) - np.sin(lat1) * np.cos(lat2) * np.cos(dlon)
    return (np.degrees(np.arctan2(y, x)) + 360) % 360

def angle_diff_np(a1, a2):
    diff = np.abs(a1 - a2) % 360
    return np.where(diff <= 180, diff, 360 - diff)


# ==========================================
# 5. CORE PROCESSOR & AUDIT REPORTING
# ==========================================

class CalculationProcessor:
    def __init__(self, cfg, cols, log_cb, cancel_token):
        self.cfg = cfg
        self.cols = cols
        self.log = log_cb
        self.cancel_token = cancel_token
        self.audit_log = []

    def log_audit(self, file_name, row_idx, column, value, reason):
        self.audit_log.append({
            "File": file_name,
            "Row": row_idx,
            "Column": column,
            "Value": str(value),
            "Reason": reason
        })

    def clean_and_validate_coords(self, df, lat_col, lon_col, file_tag):
        cleaned_lat = (
            df[lat_col].astype(str)
            .str.replace(r"[\s\xa0]", "", regex=True)
            .str.replace(",", ".", regex=False)
        )
        cleaned_lon = (
            df[lon_col].astype(str)
            .str.replace(r"[\s\xa0]", "", regex=True)
            .str.replace(",", ".", regex=False)
        )

        num_lat = pd.to_numeric(cleaned_lat, errors="coerce")
        num_lon = pd.to_numeric(cleaned_lon, errors="coerce")

        valid_series = pd.Series(True, index=df.index)

        for idx in df.index:
            row_num = idx + 2
            lat_val = num_lat[idx]
            lon_val = num_lon[idx]

            if pd.isna(lat_val) or pd.isna(lon_val):
                valid_series[idx] = False
                self.log_audit(file_tag, row_num, f"{lat_col}/{lon_col}", f"{df.loc[idx, lat_col]}/{df.loc[idx, lon_col]}", "Non-numeric or missing coordinate")
                continue

            if lat_val == 0.0 and lon_val == 0.0:
                valid_series[idx] = False
                self.log_audit(file_tag, row_num, f"{lat_col}/{lon_col}", "(0,0)", "Null Island (0,0) coordinate flagged")
                continue

            # Check for swapped lat/lon
            if (abs(lat_val) > 90 and abs(lat_val) <= 180) and abs(lon_val) <= 90:
                valid_series[idx] = False
                self.log_audit(file_tag, row_num, f"{lat_col}/{lon_col}", f"Lat:{lat_val}, Lon:{lon_val}", "Suspected Swapped Lat/Lon (Lat > 90)")
                continue

            if not (-90.0 <= lat_val <= 90.0):
                valid_series[idx] = False
                self.log_audit(file_tag, row_num, lat_col, lat_val, "Latitude out of range [-90, 90]")

            if not (-180.0 <= lon_val <= 180.0):
                valid_series[idx] = False
                self.log_audit(file_tag, row_num, lon_col, lon_val, "Longitude out of range [-180, 180]")

        df_cleaned = df.copy()
        df_cleaned["LAT_clean"] = num_lat
        df_cleaned["LON_clean"] = num_lon
        df_cleaned["_is_valid"] = valid_series
        return df_cleaned

    def run(self, df_msan_raw, df_mobile_raw):
        self.log("Validating input data and coordinates...")

        df_msan = self.clean_and_validate_coords(df_msan_raw, self.cols["msan_lat"], self.cols["msan_lon"], "MSAN")
        df_mob = self.clean_and_validate_coords(df_mobile_raw, self.cols["mob_lat"], self.cols["mob_lon"], "Mobile")

        # Handle Azimuths explicitly
        mob_azi_str = (
            df_mob[self.cols["mob_azimut"]].astype(str)
            .str.replace(r"[\s\xa0]", "", regex=True)
            .str.replace(",", ".", regex=False)
        )
        df_mob["AZI_clean"] = pd.to_numeric(mob_azi_str, errors="coerce")
        
        # Log missing azimuths
        missing_azi = df_mob["AZI_clean"].isna() & df_mob["_is_valid"]
        for idx in df_mob[missing_azi].index:
            self.log_audit("Mobile", idx + 2, self.cols["mob_azimut"], df_mob.loc[idx, self.cols["mob_azimut"]], "Missing Azimuth: Treated as Omnidirectional (Flagged)")

        # Preserve the original MSAN row index for KD-tree candidate lookup.
        df_msan_valid = df_msan[df_msan["_is_valid"]].copy()
        df_msan_valid["msan_orig_idx"] = df_msan_valid.index
        df_msan_valid = df_msan_valid.reset_index(drop=True)

        df_mob_valid = df_mob[df_mob["_is_valid"]].copy().reset_index(drop=True)

        if df_msan_valid.empty or df_mob_valid.empty:
            raise ValueError(I18N[self.cfg["lang"]]["err_no_valid_points"])

        self.log("Building Spatial Index (3D KD-Tree)...")
        unique_sites = (
            df_mob_valid.groupby(self.cols["mob_site"])[["LAT_clean", "LON_clean"]]
            .mean()
            .reset_index()
        )

        site_lat_rad = np.radians(unique_sites["LAT_clean"].values)
        site_lon_rad = np.radians(unique_sites["LON_clean"].values)
        site_x = np.cos(site_lat_rad) * np.cos(site_lon_rad)
        site_y = np.cos(site_lat_rad) * np.sin(site_lon_rad)
        site_z = np.sin(site_lat_rad)

        tree = cKDTree(np.column_stack([site_x, site_y, site_z]))

        m_lat_rad = np.radians(df_msan_valid["LAT_clean"].values)
        m_lon_rad = np.radians(df_msan_valid["LON_clean"].values)
        m_x = np.cos(m_lat_rad) * np.cos(m_lon_rad)
        m_y = np.cos(m_lat_rad) * np.sin(m_lon_rad)
        m_z = np.sin(m_lat_rad)

        k_pool = min(self.cfg["cand_pool"], len(unique_sites))
        _, indices_matrix = tree.query(np.column_stack([m_x, m_y, m_z]), k=k_pool)

        # Pre-group cell sectors by site
        cells_by_site = {}
        for site_name, grp in df_mob_valid.groupby(self.cols["mob_site"]):
            cells_by_site[site_name] = {
                "cells": grp[self.cols["mob_cell"]].values,
                "azimuths": grp["AZI_clean"].values,
            }

        self.log("Calculating nearest distance and best sector alignments...")
        results = []
        top_n = self.cfg["top_n"]
        max_dist_cap = self.cfg["max_dist_km"]

        for idx, row in df_msan.iterrows():
            if self.cancel_token.is_set():
                self.log("Processing cancelled by user.")
                return None, None

            out_row = {
                "Ligne_CSV": idx + 2,
                "Site_MSAN": row[self.cols["msan_site"]],
                "MSAN_LON": row[self.cols["msan_lon"]],
                "MSAN_LAT": row[self.cols["msan_lat"]],
                "Statut": "OK",
            }

            if not row["_is_valid"]:
                out_row["Statut"] = "Ignoré: Coordonnées Invalides"
                for i in range(1, top_n + 1):
                    out_row[f"Top{i}_Site_Mobile"] = None
                    out_row[f"Top{i}_Mobile_LAT"] = None
                    out_row[f"Top{i}_Mobile_LON"] = None
                    out_row[f"Top{i}_Distance_km"] = None
                    out_row[f"Top{i}_Best_Cell"] = None
                    out_row[f"Top{i}_Azimut"] = None
                    out_row[f"Top{i}_Angle_Dev"] = None
                results.append(out_row)
                continue

            # Lookup KDTree pre-computed candidates
            m_valid_pos = df_msan_valid[df_msan_valid["msan_orig_idx"] == idx].index[0] if "msan_orig_idx" in df_msan_valid else None
            
            # Direct calculation for candidates
            m_lat, m_lon = row["LAT_clean"], row["LON_clean"]
            
            cand_indices = indices_matrix[m_valid_pos] if m_valid_pos is not None else range(len(unique_sites))
            if isinstance(cand_indices, np.int64):
                cand_indices = [cand_indices]

            cand_sites = unique_sites.iloc[cand_indices]

            dists = haversine_np(m_lat, m_lon, cand_sites["LAT_clean"].values, cand_sites["LON_clean"].values)
            bearings = bearing_np(cand_sites["LAT_clean"].values, cand_sites["LON_clean"].values, m_lat, m_lon)

            evals = []
            for i, (_, c_row) in enumerate(cand_sites.iterrows()):
                s_name = c_row[self.cols["mob_site"]]
                d_km = dists[i]
                brg = bearings[i]

                if d_km > max_dist_cap:
                    continue

                if s_name in cells_by_site:
                    c_data = cells_by_site[s_name]
                    azimuths = c_data["azimuths"]
                    
                    # If azimuth is NaN, default dev to 0 (Omnidirectional)
                    valid_azis = np.nan_to_num(azimuths, nan=brg)
                    devs = angle_diff_np(valid_azis, brg)
                    best_i = np.argmin(devs)

                    evals.append({
                        "site": s_name,
                        "lat": float(c_row["LAT_clean"]),
                        "lon": float(c_row["LON_clean"]),
                        "dist": round(d_km, 3),
                        "cell": c_data["cells"][best_i],
                        "azimut": "Omni" if np.isnan(azimuths[best_i]) else azimuths[best_i],
                        "dev": round(devs[best_i], 1)
                    })

            evals.sort(key=lambda x: x["dist"])

            for i in range(1, top_n + 1):
                if (i - 1) < len(evals):
                    ev = evals[i - 1]
                    out_row[f"Top{i}_Site_Mobile"] = ev["site"]
                    out_row[f"Top{i}_Mobile_LAT"] = ev["lat"]
                    out_row[f"Top{i}_Mobile_LON"] = ev["lon"]
                    out_row[f"Top{i}_Distance_km"] = ev["dist"]
                    out_row[f"Top{i}_Best_Cell"] = ev["cell"]
                    out_row[f"Top{i}_Azimut"] = ev["azimut"]
                    out_row[f"Top{i}_Angle_Dev"] = ev["dev"]
                else:
                    out_row[f"Top{i}_Site_Mobile"] = None
                    out_row[f"Top{i}_Mobile_LAT"] = None
                    out_row[f"Top{i}_Mobile_LON"] = None
                    out_row[f"Top{i}_Distance_km"] = None
                    out_row[f"Top{i}_Best_Cell"] = None
                    out_row[f"Top{i}_Azimut"] = None
                    out_row[f"Top{i}_Angle_Dev"] = None

            results.append(out_row)

        res_df = pd.DataFrame(results)
        audit_df = pd.DataFrame(self.audit_log)
        return res_df, audit_df


# ==========================================
# 6. EXCEL FORMATTER & EXPORTER
# ==========================================

class ExcelExporter:
    @staticmethod
    def export(results_df, audit_df, filepath, decimal_sep=",", kml_audit_df=None):
        if not OPENPYXL_AVAILABLE:
            # Fallback to CSV if openpyxl is not installed
            csv_file = filepath.replace(".xlsx", ".csv")
            results_df.to_csv(csv_file, sep=";", decimal=decimal_sep, index=False, encoding="utf-8-sig")
            return csv_file

        wb = openpyxl.Workbook()
        ws_res = wb.active
        ws_res.title = "MSAN_Mobile_Distances"

        # Apply decimal separator formatting if string conversion required
        df_export = results_df.copy()
        if decimal_sep == ",":
            float_cols = df_export.select_dtypes(include=["float", "float64"]).columns
            for c in float_cols:
                df_export[c] = df_export[c].astype(str).str.replace(".", ",", regex=False)

        # Style Definitions
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for r in dataframe_to_rows(df_export, index=False, header=True):
            ws_res.append(r)

        for cell in ws_res[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws_res.iter_rows(min_row=2, max_row=ws_res.max_row, max_col=ws_res.max_column):
            for cell in row:
                cell.border = thin_border

        # Validation Audit Sheet
        ws_audit = wb.create_sheet(title="Validation_Report")
        if not audit_df.empty:
            for r in dataframe_to_rows(audit_df, index=False, header=True):
                ws_audit.append(r)
            for cell in ws_audit[1]:
                cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        # KML Generation Audit Sheet
        if kml_audit_df is not None:
            ws_kml = wb.create_sheet(title="KML_Generation_Report")
            if not kml_audit_df.empty:
                for r in dataframe_to_rows(kml_audit_df, index=False, header=True):
                    ws_kml.append(r)
                for cell in ws_kml[1]:
                    cell.fill = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
                    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                for col in ws_kml.columns:
                    ws_kml.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 2, 60)

        wb.save(filepath)
        return filepath


# ==========================================
# 7. MAIN APPLICATION GUI
# ==========================================

class ApplicationGUI:
    def __init__(self, root):
        self.root = root
        self.lang = "FR"
        self.cancel_token = threading.Event()
        self.msg_queue = queue.Queue()

        self.df_msan_raw = None
        self.df_mob_raw = None

        self.load_app_config()
        self.setup_ui()
        self.apply_translations()

        self.root.after(100, self.process_queue)

    def load_app_config(self):
        self.app_cfg = {
            "last_msan_path": "",
            "last_mob_path": "",
            "output_dir": "",
            "top_n": 3,
            "cand_pool": 15,
            "max_dist_km": 50.0,
            "locale_sep": ",",
            "export_kml": True,
            "window_size": "920x720"
        }
        if os.path.exists(CONFIG_FILE):
            try:
                with open(CONFIG_FILE, "r") as f:
                    self.app_cfg.update(json.load(f))
            except Exception:
                pass

    def save_app_config(self):
        try:
            self.app_cfg["window_size"] = self.root.geometry()
            with open(CONFIG_FILE, "w") as f:
                json.dump(self.app_cfg, f, indent=4)
        except Exception:
            pass

    def _build_logo(self, parent):
        """Logo: logo.png si présent, sinon badge dessiné (aucun fichier requis)."""
        logo_path = resource_path(LOGO_FILE)
        if os.path.exists(logo_path):
            try:
                img = tk.PhotoImage(file=logo_path)
                factor = max(1, img.height() // 40)
                if factor > 1:
                    img = img.subsample(factor, factor)
                self._logo_img = img  # garder une référence (sinon effacé par le GC)
                return tk.Label(parent, image=img)
            except tk.TclError:
                pass
        canvas = tk.Canvas(parent, width=40, height=40, highlightthickness=0)
        canvas.create_oval(2, 2, 38, 38, fill="#1F4E78", outline="#0B2A45", width=2)
        canvas.create_text(20, 20, text="19W", fill="white", font=("Segoe UI", 10, "bold"))
        return canvas

    def setup_ui(self):
        self.i18n_labels = []  # (widget, i18n_key, strip_colon)
        self.root.geometry(self.app_cfg.get("window_size", "920x720"))
        self.root.minsize(800, 600)

        # Top Bar (Language & Title)
        top_bar = ttk.Frame(self.root)
        top_bar.pack(fill="x", padx=15, pady=8)

        self.logo_widget = self._build_logo(top_bar)
        self.logo_widget.pack(side="left", padx=(0, 8))

        self.lbl_title = tk.Label(top_bar, text="", font=("Segoe UI", 14, "bold"), fg="#1F4E78")
        self.lbl_title.pack(side="left")

        self.combo_lang = ttk.Combobox(top_bar, values=["FR", "EN"], width=5, state="readonly")
        self.combo_lang.set(self.lang)
        self.combo_lang.pack(side="right")
        self.combo_lang.bind("<<ComboboxSelected>>", self.on_lang_change)

        # Main Notebook
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=15, pady=5)

        self.tab_workflow = ttk.Frame(self.notebook)
        self.tab_settings = ttk.Frame(self.notebook)
        self.tab_audit = ttk.Frame(self.notebook)
        self.tab_kml_audit = ttk.Frame(self.notebook)

        self.notebook.add(self.tab_workflow, text="")
        self.notebook.add(self.tab_settings, text="")
        self.notebook.add(self.tab_audit, text="")
        self.notebook.add(self.tab_kml_audit, text="")

        # --- TAB 1: WORKFLOW ---
        self.sec_files = ttk.LabelFrame(self.tab_workflow, text="")
        self.sec_files.pack(fill="x", padx=10, pady=5)

        lbl = ttk.Label(self.sec_files, text="MSAN File:")
        lbl.grid(row=0, column=0, sticky="w", padx=8, pady=5)
        self.i18n_labels.append((lbl, "msan_file", False))
        self.ent_msan = ttk.Entry(self.sec_files, width=60)
        self.ent_msan.insert(0, self.app_cfg.get("last_msan_path", ""))
        self.ent_msan.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        self.btn_browse_msan = ttk.Button(self.sec_files, text="...", command=lambda: self.browse_file("msan"))
        self.btn_browse_msan.grid(row=0, column=2, padx=8, pady=5)

        lbl = ttk.Label(self.sec_files, text="Mobile File:")
        lbl.grid(row=1, column=0, sticky="w", padx=8, pady=5)
        self.i18n_labels.append((lbl, "mob_file", False))
        self.ent_mob = ttk.Entry(self.sec_files, width=60)
        self.ent_mob.insert(0, self.app_cfg.get("last_mob_path", ""))
        self.ent_mob.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        self.btn_browse_mob = ttk.Button(self.sec_files, text="...", command=lambda: self.browse_file("mob"))
        self.btn_browse_mob.grid(row=1, column=2, padx=8, pady=5)

        self.lbl_out_dir = ttk.Label(self.sec_files, text="Output Folder:")
        self.lbl_out_dir.grid(row=2, column=0, sticky="w", padx=8, pady=5)
        self.ent_out = ttk.Entry(self.sec_files, width=60)
        self.ent_out.insert(0, self.app_cfg.get("output_dir", ""))
        self.ent_out.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        self.btn_browse_out = ttk.Button(self.sec_files, text="...", command=self.browse_output_dir)
        self.btn_browse_out.grid(row=2, column=2, padx=8, pady=5)

        self.sec_files.columnconfigure(1, weight=1)

        # Action Buttons Frame
        btn_bar = ttk.Frame(self.tab_workflow)
        btn_bar.pack(fill="x", padx=10, pady=5)

        self.btn_discovery = tk.Button(btn_bar, text="", bg="#007BFF", fg="white", font=("Segoe UI", 9, "bold"), command=self.run_discovery)
        self.btn_discovery.pack(side="left", padx=5)

        self.btn_load_prof = ttk.Button(btn_bar, text="", command=self.load_profile)
        self.btn_load_prof.pack(side="left", padx=5)

        self.btn_save_prof = ttk.Button(btn_bar, text="", command=self.save_profile)
        self.btn_save_prof.pack(side="left", padx=5)

        # Mapping Section
        self.sec_mapping = ttk.LabelFrame(self.tab_workflow, text="")
        self.sec_mapping.pack(fill="x", padx=10, pady=5)

        map_left = ttk.Frame(self.sec_mapping)
        map_left.pack(side="left", fill="both", expand=True, padx=5, pady=5)
        map_right = ttk.Frame(self.sec_mapping)
        map_right.pack(side="right", fill="both", expand=True, padx=5, pady=5)

        self.combos = {}
        self.msan_fields = [("lbl_site_name", "msan_site", ["site", "nom"]), ("lbl_lon", "msan_lon", ["lon"]), ("lbl_lat", "msan_lat", ["lat"])]
        self.mob_fields = [("lbl_site_name", "mob_site", ["site"]), ("lbl_cell_name", "mob_cell", ["cell"]), ("lbl_azimut", "mob_azimut", ["azi"]), ("lbl_lon", "mob_lon", ["lon"]), ("lbl_lat", "mob_lat", ["lat"])]

        for i, (lbl_key, key, _) in enumerate(self.msan_fields):
            lbl = ttk.Label(map_left, text=lbl_key)
            lbl.grid(row=i, column=0, sticky="w", padx=5, pady=3)
            self.i18n_labels.append((lbl, lbl_key, False))
            cb = ttk.Combobox(map_left, state="disabled", width=22)
            cb.grid(row=i, column=1, padx=5, pady=3)
            self.combos[key] = cb

        for i, (lbl_key, key, _) in enumerate(self.mob_fields):
            lbl = ttk.Label(map_right, text=lbl_key)
            lbl.grid(row=i, column=0, sticky="w", padx=5, pady=3)
            self.i18n_labels.append((lbl, lbl_key, False))
            cb = ttk.Combobox(map_right, state="disabled", width=22)
            cb.grid(row=i, column=1, padx=5, pady=3)
            self.combos[key] = cb

        # Execution Controls
        exec_bar = ttk.Frame(self.tab_workflow)
        exec_bar.pack(fill="x", padx=10, pady=10)

        self.btn_run = tk.Button(exec_bar, text="", bg="#28A745", fg="white", font=("Segoe UI", 10, "bold"), state="disabled", command=self.start_processing)
        self.btn_run.pack(side="left", fill="x", expand=True, padx=5)

        self.btn_cancel = tk.Button(exec_bar, text="", bg="#DC3545", fg="white", font=("Segoe UI", 10, "bold"), state="disabled", command=self.cancel_processing)
        self.btn_cancel.pack(side="right", padx=5)

        self.progress = ttk.Progressbar(self.tab_workflow, mode="indeterminate")
        self.progress.pack(fill="x", padx=15, pady=5)

        # Log Text Box
        self.log_text = tk.Text(self.tab_workflow, height=8, state="disabled", bg="#F8F9FA", font=("Consolas", 9))
        self.log_text.pack(fill="both", expand=True, padx=15, pady=5)

        # --- TAB 2: SETTINGS ---
        self.sec_engine_cfg = ttk.LabelFrame(self.tab_settings, text="")
        self.sec_engine_cfg.pack(fill="both", expand=True, padx=15, pady=15)

        lbl = ttk.Label(self.sec_engine_cfg, text="Top N Neighbors:")
        lbl.grid(row=0, column=0, sticky="w", padx=10, pady=10)
        self.i18n_labels.append((lbl, "lbl_top_n", False))
        self.spin_top_n = ttk.Spinbox(self.sec_engine_cfg, from_=1, to=10, width=8)
        self.spin_top_n.set(self.app_cfg["top_n"])
        self.spin_top_n.grid(row=0, column=1, padx=10, pady=10, sticky="w")

        lbl = ttk.Label(self.sec_engine_cfg, text="Candidate Pool (k_search):")
        lbl.grid(row=1, column=0, sticky="w", padx=10, pady=10)
        self.i18n_labels.append((lbl, "lbl_cand_pool", False))
        self.spin_cand_pool = ttk.Spinbox(self.sec_engine_cfg, from_=5, to=50, width=8)
        self.spin_cand_pool.set(self.app_cfg["cand_pool"])
        self.spin_cand_pool.grid(row=1, column=1, padx=10, pady=10, sticky="w")

        lbl = ttk.Label(self.sec_engine_cfg, text="Max Distance Cap (km):")
        lbl.grid(row=2, column=0, sticky="w", padx=10, pady=10)
        self.i18n_labels.append((lbl, "lbl_max_dist", False))
        self.ent_max_dist = ttk.Entry(self.sec_engine_cfg, width=10)
        self.ent_max_dist.insert(0, str(self.app_cfg["max_dist_km"]))
        self.ent_max_dist.grid(row=2, column=1, padx=10, pady=10, sticky="w")

        lbl = ttk.Label(self.sec_engine_cfg, text="Decimal Separator:")
        lbl.grid(row=3, column=0, sticky="w", padx=10, pady=10)
        self.i18n_labels.append((lbl, "lbl_locale", False))
        self.combo_locale = ttk.Combobox(self.sec_engine_cfg, values=[", (Comma)", ". (Dot)"], width=12, state="readonly")
        self.combo_locale.set(", (Comma)" if self.app_cfg["locale_sep"] == "," else ". (Dot)")
        self.combo_locale.grid(row=3, column=1, padx=10, pady=10, sticky="w")

        self.var_kml = tk.BooleanVar(value=self.app_cfg["export_kml"])
        self.chk_kml = ttk.Checkbutton(self.sec_engine_cfg, text="Export KML Map File", variable=self.var_kml)
        self.chk_kml.grid(row=4, column=0, columnspan=2, sticky="w", padx=10, pady=10)
        self.i18n_labels.append((self.chk_kml, "lbl_kml_export", True))

        # --- TAB 3: AUDIT TABLE ---
        self.tree_audit = ttk.Treeview(self.tab_audit, columns=("File", "Row", "Column", "Value", "Reason"), show="headings")
        for c in ("File", "Row", "Column", "Value", "Reason"):
            self.tree_audit.heading(c, text=c)
            self.tree_audit.column(c, width=120)
        self.tree_audit.pack(fill="both", expand=True, padx=10, pady=10)

        # --- TAB 4: KML GENERATION LOG ---
        self.tree_kml_audit = ttk.Treeview(
            self.tab_kml_audit,
            columns=("Layer", "Status", "Object", "Name", "Source Row", "Reason", "Details"),
            show="headings"
        )
        for c in ("Layer", "Status", "Object", "Name", "Source Row", "Reason", "Details"):
            self.tree_kml_audit.heading(c, text=c)
            self.tree_kml_audit.column(c, width=150)
        self.tree_kml_audit.column("Details", width=300)
        self.tree_kml_audit.column("Reason", width=280)
        self.tree_kml_audit.pack(fill="both", expand=True, padx=10, pady=10)

        # Status Footer
        footer = ttk.Frame(self.root)
        footer.pack(side="bottom", fill="x", padx=15, pady=5)
        self.lbl_status = ttk.Label(footer, text="", font=("Segoe UI", 9, "italic"))
        self.lbl_status.pack(side="left", fill="x", expand=True)
        self.lbl_signature = tk.Label(footer, text=f"© {SIGNATURE}", font=("Segoe UI", 9, "bold"), fg="#1F4E78")
        self.lbl_signature.pack(side="right")

    def on_lang_change(self, event=None):
        self.lang = self.combo_lang.get()
        self.apply_translations()

    def apply_translations(self):
        t = I18N[self.lang]
        self.root.title(t["title"])
        self.lbl_title.config(text=t["title"])

        self.notebook.tab(self.tab_workflow, text=t["tab_workflow"])
        self.notebook.tab(self.tab_settings, text=t["tab_settings"])
        self.notebook.tab(self.tab_audit, text=t["tab_audit"])
        self.notebook.tab(self.tab_kml_audit, text=" Rapport KML " if self.lang == "FR" else " KML Report ")

        self.sec_files.config(text=t["sec_files"])
        self.sec_mapping.config(text=t["sec_mapping"])
        self.sec_engine_cfg.config(text=t["sec_engine_cfg"])
        self.lbl_out_dir.config(text=t["lbl_out_dir"])

        self.btn_discovery.config(text=t["btn_discovery"])
        self.btn_load_prof.config(text=t["btn_load_profile"])
        self.btn_save_prof.config(text=t["btn_save_profile"])
        self.btn_run.config(text=t["btn_run"])
        self.btn_cancel.config(text=t["btn_cancel"])

        for widget, key, strip_colon in self.i18n_labels:
            txt = t[key]
            widget.config(text=txt.rstrip(" :") if strip_colon else txt)

        self.lbl_status.config(text=t["lbl_status_init"])

    def browse_file(self, file_type):
        path = filedialog.askopenfilename(filetypes=[("Data Files", "*.csv;*.xlsx;*.xls;*.xlsb"), ("All Files", "*.*")])
        if path:
            if file_type == "msan":
                self.ent_msan.delete(0, tk.END)
                self.ent_msan.insert(0, path)
            else:
                self.ent_mob.delete(0, tk.END)
                self.ent_mob.insert(0, path)

    def browse_output_dir(self):
        initial = self.ent_out.get().strip()
        if not initial or not os.path.isdir(initial):
            initial = os.path.dirname(self.ent_msan.get().strip()) or os.getcwd()
        path = filedialog.askdirectory(initialdir=initial)
        if path:
            path = os.path.normpath(path)
            self.ent_out.delete(0, tk.END)
            self.ent_out.insert(0, path)
            self.app_cfg["output_dir"] = path
            self.save_app_config()

    def log(self, text):
        self.msg_queue.put(("LOG", text))

    def process_queue(self):
        try:
            while True:
                msg_type, content = self.msg_queue.get_nowait()
                if msg_type == "LOG":
                    self.log_text.config(state="normal")
                    self.log_text.insert(tk.END, f"[{datetime.now().strftime('%H:%M:%S')}] {content}\n")
                    self.log_text.see(tk.END)
                    self.log_text.config(state="disabled")
                elif msg_type == "STATUS":
                    self.lbl_status.config(text=content)
        except queue.Empty:
            pass
        self.root.after(100, self.process_queue)

    def run_discovery(self):
        msan_path = self.ent_msan.get().strip()
        mob_path = self.ent_mob.get().strip()

        if not msan_path or not mob_path:
            messagebox.showwarning("Warning", I18N[self.lang]["warn_no_files"])
            return

        try:
            self.df_msan_raw, msan_desc = DataIngestionEngine.read_file(msan_path)
            self.df_mob_raw, mob_desc = DataIngestionEngine.read_file(mob_path)

            self.log(f"Loaded MSAN File: {msan_desc} ({len(self.df_msan_raw)} rows)")
            self.log(f"Loaded Mobile File: {mob_desc} ({len(self.df_mob_raw)} rows)")

            msan_cols = list(self.df_msan_raw.columns)
            mob_cols = list(self.df_mob_raw.columns)

            for _, key, hints in self.msan_fields:
                cb = self.combos[key]
                cb.config(state="readonly", values=msan_cols)
                cb.set(next((c for c in msan_cols if any(h in c.lower() for h in hints)), msan_cols[0]))

            for _, key, hints in self.mob_fields:
                cb = self.combos[key]
                cb.config(state="readonly", values=mob_cols)
                cb.set(next((c for c in mob_cols if any(h in c.lower() for h in hints)), mob_cols[0]))

            self.btn_run.config(state="normal")
            self.lbl_status.config(text=I18N[self.lang]["lbl_status_discovered"])

            # Save paths in app config
            self.app_cfg["last_msan_path"] = msan_path
            self.app_cfg["last_mob_path"] = mob_path
            self.save_app_config()

        except Exception as e:
            self.log(f"Error reading files: {str(e)}")
            messagebox.showerror("Error", f"Failed to ingest files:\n{str(e)}")

    def save_profile(self):
        path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("Mapping Profile", "*.json")])
        if path:
            mapping = {k: v.get() for k, v in self.combos.items()}
            with open(path, "w") as f:
                json.dump(mapping, f, indent=4)
            self.log(f"Saved mapping profile to {path}")

    def load_profile(self):
        path = filedialog.askopenfilename(filetypes=[("Mapping Profile", "*.json")])
        if path:
            try:
                with open(path, "r") as f:
                    mapping = json.load(f)
                for k, v in mapping.items():
                    if k in self.combos:
                        self.combos[k].set(v)
                self.log(f"Loaded mapping profile from {path}")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to load profile: {str(e)}")

    def start_processing(self):
        # Dossier de sortie : vide = dossier de travail courant (comportement historique)
        out_dir = self.ent_out.get().strip() or os.getcwd()
        try:
            os.makedirs(out_dir, exist_ok=True)
            if not os.access(out_dir, os.W_OK):
                raise PermissionError("Write access denied")
        except Exception as e:
            messagebox.showerror("Error", f"{I18N[self.lang]['err_out_dir']}\n{out_dir}\n{e}")
            return
        self.app_cfg["output_dir"] = self.ent_out.get().strip()
        self.save_app_config()

        self.cancel_token.clear()
        self.btn_run.config(state="disabled")
        self.btn_cancel.config(state="normal")
        self.progress.start()

        # Build runtime config
        cfg = {
            "lang": self.lang,
            "top_n": int(self.spin_top_n.get()),
            "cand_pool": int(self.spin_cand_pool.get()),
            "max_dist_km": float(self.ent_max_dist.get()),
            "locale_sep": "," if "," in self.combo_locale.get() else ".",
            "export_kml": self.var_kml.get()
        }
        cols = {k: v.get() for k, v in self.combos.items()}

        def worker():
            try:
                processor = CalculationProcessor(cfg, cols, self.log, self.cancel_token)
                res_df, audit_df = processor.run(self.df_msan_raw, self.df_mob_raw)

                if res_df is None:
                    self.msg_queue.put(("STATUS", I18N[self.lang]["lbl_status_cancelled"]))
                    return

                # Populate Validation Audit Table UI
                for item in self.tree_audit.get_children():
                    self.tree_audit.delete(item)
                for _, row in audit_df.iterrows():
                    self.tree_audit.insert("", "end", values=(row["File"], row["Row"], row["Column"], row["Value"], row["Reason"]))

                # Reset KML Generation Log for this run
                for item in self.tree_kml_audit.get_children():
                    self.tree_kml_audit.delete(item)

                # Export Results
                ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                out_excel = os.path.join(out_dir, f"MSAN_Mobile_Results_{ts}.xlsx")
                kml_audit_df = pd.DataFrame()
                if cfg["export_kml"]:
                    out_kml = os.path.join(out_dir, f"MSAN_Mobile_Map_{ts}.kml")
                    kml_audit_df = KMLExporter.export(
                        res_df,
                        out_kml,
                        mobile_df=self.df_mob_raw,
                        mobile_cols=cols
                    )
                    for _, row in kml_audit_df.iterrows():
                        self.tree_kml_audit.insert(
                            "", "end",
                            values=(
                                row["Layer"], row["Status"], row["Object"],
                                row["Name"], row["Source Row"],
                                row["Reason"], row["Details"]
                            )
                        )
                    self.log(f"KML generation report created: {len(kml_audit_df)} entries")
                    self.log(f"Exported KML Spatial Map: {out_kml}")

                saved_path = ExcelExporter.export(
                    res_df, audit_df, out_excel,
                    decimal_sep=cfg["locale_sep"],
                    kml_audit_df=kml_audit_df
                )
                self.log(f"Exported Excel Results: {saved_path}")

                self.msg_queue.put(("STATUS", I18N[self.lang]["lbl_status_success"]))
                messagebox.showinfo("Success", f"Processing complete!\nGenerated: {saved_path}")

            except Exception as e:
                self.log(f"Execution Error: {str(e)}")
                self.msg_queue.put(("STATUS", I18N[self.lang]["lbl_status_error"]))
                messagebox.showerror("Execution Error", str(e))
            finally:
                self.progress.stop()
                self.btn_run.config(state="normal")
                self.btn_cancel.config(state="disabled")

        threading.Thread(target=worker, daemon=True).start()

    def cancel_processing(self):
        self.cancel_token.set()
        self.log("Cancelling operation...")


# ==========================================
# 8. UNIT TESTS AGAINST KNOWN REFERENCE
# ==========================================

class TestSpatialEngine(unittest.TestCase):
    def test_haversine_distance(self):
        # Paris to London approx ~343 km
        paris_lat, paris_lon = 48.8566, 2.3522
        london_lat, london_lon = 51.5074, -0.1278
        dist = haversine_np(paris_lat, paris_lon, london_lat, london_lon)
        self.assertAlmostEqual(dist, 343.5, delta=5.0)

    def test_bearing_calculation(self):
        # Due North
        brg = bearing_np(0.0, 0.0, 1.0, 0.0)
        self.assertAlmostEqual(brg, 0.0, delta=0.1)

    def test_angle_difference(self):
        diff = angle_diff_np(np.array([10]), np.array([350]))
        self.assertEqual(diff[0], 20)


# ==========================================
# MAIN ENTRY POINT
# ==========================================

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--test":
        unittest.main(argv=[sys.argv[0]])
    else:
        root = tk.Tk()
        app = ApplicationGUI(root)
        root.protocol("WM_DELETE_WINDOW", lambda: (app.save_app_config(), root.destroy()))
        root.mainloop()
